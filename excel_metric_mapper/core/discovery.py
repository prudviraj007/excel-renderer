"""Discovery engine for finding month headers and metric labels."""
from typing import List, Dict, Optional, Tuple
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .models import MonthHeaderCandidate, MetricLabelCandidate, SourceMetric
from .month_parser import normalize_month, is_likely_month_value
from rapidfuzz import fuzz
import re


def find_month_headers(
    workbook_path: str,
    sheet_name: str,
    target_month: Optional[str] = None,
    max_rows: int = 50,
    max_cols: int = 50
) -> List[MonthHeaderCandidate]:
    """
    Find potential month header cells in a worksheet.
    
    Args:
        workbook_path: Path to workbook
        sheet_name: Sheet name
        target_month: Optional specific month to find (YYYY-MM)
        max_rows: Max rows to scan
        max_cols: Max columns to scan
    
    Returns:
        List of MonthHeaderCandidate objects
    """
    candidates = []
    
    try:
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet = wb[sheet_name]
        
        # Scan cells for month values
        for row in range(1, min(sheet.max_row + 1, max_rows + 1)):
            for col in range(1, min(sheet.max_column + 1, max_cols + 1)):
                cell = sheet.cell(row=row, column=col)
                value = cell.value
                
                if value is None or value == "":
                    continue
                
                # Check if this looks like a month
                if not is_likely_month_value(value):
                    continue
                
                # Try to normalize
                normalized = normalize_month(value)
                if not normalized:
                    continue
                
                # Filter by target month if specified
                if target_month and normalized != target_month:
                    continue
                
                # Determine orientation based on surrounding cells
                orientation = _detect_month_orientation(sheet, row, col)
                
                # Calculate confidence
                confidence = _calculate_month_confidence(sheet, row, col, normalized, orientation)
                
                # Build evidence string
                evidence = _build_month_evidence(sheet, row, col, normalized, orientation)
                
                candidate = MonthHeaderCandidate(
                    sheet_name=sheet_name,
                    cell_address=cell.coordinate,
                    row=row,
                    col=col,
                    raw_value=value,
                    normalized_month=normalized,
                    confidence=confidence,
                    orientation=orientation,
                    evidence=evidence
                )
                
                candidates.append(candidate)
        
        wb.close()
        
    except Exception as e:
        print(f"Error finding month headers in {sheet_name}: {str(e)}")
    
    return candidates


def _count_months_in_row(sheet, row: int, span: int = 24) -> int:
    """Count month-like headers on a row (typical column header layout)."""
    count = 0
    for c in range(1, min(sheet.max_column + 1, span + 1)):
        value = sheet.cell(row=row, column=c).value
        if is_likely_month_value(value) and normalize_month(value):
            count += 1
    return count


def _count_months_in_column(sheet, col: int, span: int = 24) -> int:
    """Count month-like headers in a column (typical row header layout)."""
    count = 0
    for r in range(1, min(sheet.max_row + 1, span + 1)):
        value = sheet.cell(row=r, column=col).value
        if is_likely_month_value(value) and normalize_month(value):
            count += 1
    return count


def _detect_month_orientation(sheet, row: int, col: int) -> str:
    """
    Detect if months are arranged horizontally (in columns) or vertically (in rows).

    Uses the full header row/column, not only neighbors. This avoids misclassifying
    a lone target month (e.g. Oct-26 at K4) as vertical when the row is Jan–Oct.
    """
    row_months = _count_months_in_row(sheet, row)
    col_months = _count_months_in_column(sheet, col)

    if row_months >= 2 and row_months >= col_months:
        return "horizontal"
    if col_months >= 2 and col_months > row_months:
        return "vertical"

    # Single month visible (e.g. filtered to target month only): infer from data shape
    if _check_numeric_data_nearby(sheet, row, col, "horizontal"):
        return "horizontal"
    if _check_numeric_data_nearby(sheet, row, col, "vertical"):
        return "vertical"

    # Typical management reports: months across the top
    return "horizontal"


def _calculate_month_confidence(sheet, row: int, col: int, normalized_month: str, orientation: str) -> float:
    """Calculate confidence score for a month header candidate."""
    confidence = 0.5  # Base confidence
    
    # Check for sequential months nearby
    has_sequence = _check_month_sequence(sheet, row, col, orientation)
    if has_sequence:
        confidence += 0.3
    
    # Check if in first few rows/columns (typical for headers)
    if orientation == 'horizontal' and row <= 5:
        confidence += 0.1
    elif orientation == 'vertical' and col <= 3:
        confidence += 0.1
    
    # Check if there are numeric values nearby (indicates data region)
    has_data_nearby = _check_numeric_data_nearby(sheet, row, col, orientation)
    if has_data_nearby:
        confidence += 0.1
    
    return min(confidence, 1.0)


def _check_month_sequence(sheet, row: int, col: int, orientation: str) -> bool:
    """Check if there's a sequence of months nearby."""
    month_count = 0
    
    if orientation == 'horizontal':
        # Check cells to the right
        for c in range(col, min(col + 5, sheet.max_column + 1)):
            cell = sheet.cell(row=row, column=c)
            if is_likely_month_value(cell.value):
                month_count += 1
    else:
        # Check cells below
        for r in range(row, min(row + 5, sheet.max_row + 1)):
            cell = sheet.cell(row=r, column=col)
            if is_likely_month_value(cell.value):
                month_count += 1
    
    return month_count >= 3


def _check_numeric_data_nearby(sheet, row: int, col: int, orientation: str) -> bool:
    """Check if there are numeric values nearby (indicates data region)."""
    if orientation == 'horizontal':
        # Check cells below
        for r in range(row + 1, min(row + 6, sheet.max_row + 1)):
            cell = sheet.cell(row=r, column=col)
            if isinstance(cell.value, (int, float)):
                return True
    else:
        # Check cells to the right
        for c in range(col + 1, min(col + 6, sheet.max_column + 1)):
            cell = sheet.cell(row=row, column=c)
            if isinstance(cell.value, (int, float)):
                return True
    
    return False


def _build_month_evidence(sheet, row: int, col: int, normalized_month: str, orientation: str) -> str:
    """Build evidence string for month candidate."""
    parts = []
    parts.append(f"Normalized to {normalized_month}")
    parts.append(f"{orientation} layout")
    
    if _check_month_sequence(sheet, row, col, orientation):
        parts.append("part of month sequence")
    
    if _check_numeric_data_nearby(sheet, row, col, orientation):
        parts.append("numeric data nearby")
    
    return "; ".join(parts)


def find_metric_labels(
    workbook_path: str,
    sheet_name: str,
    source_metrics: List[SourceMetric],
    max_rows: int = 100,
    max_cols: int = 20
) -> List[MetricLabelCandidate]:
    """
    Find potential metric label cells in a worksheet.
    
    Args:
        workbook_path: Path to workbook
        sheet_name: Sheet name
        source_metrics: List of source metrics to match
        max_rows: Max rows to scan
        max_cols: Max columns to scan
    
    Returns:
        List of MetricLabelCandidate objects
    """
    candidates = []
    
    try:
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet = wb[sheet_name]
        
        # Build lookup for metrics
        metric_by_id = {m.metric_id: m for m in source_metrics}
        metric_by_name = {m.metric_name.lower(): m for m in source_metrics}
        
        # Scan cells for text that might be metric labels
        for row in range(1, min(sheet.max_row + 1, max_rows + 1)):
            for col in range(1, min(sheet.max_column + 1, max_cols + 1)):
                cell = sheet.cell(row=row, column=col)
                value = cell.value
                
                if not value or not isinstance(value, str):
                    continue
                
                label_text = value.strip()
                if not label_text or len(label_text) < 2:
                    continue
                
                # Skip if it looks like a month
                if is_likely_month_value(value):
                    continue
                
                # Try to match against source metrics
                best_match = _find_best_metric_match(label_text, source_metrics)
                
                if best_match:
                    metric, match_type, score = best_match
                    
                    candidate = MetricLabelCandidate(
                        sheet_name=sheet_name,
                        cell_address=cell.coordinate,
                        row=row,
                        col=col,
                        label_text=label_text,
                        matched_metric_id=metric.metric_id,
                        matched_metric_name=metric.metric_name,
                        match_score=score,
                        match_type=match_type,
                        evidence=f"{match_type}: '{label_text}' → '{metric.metric_name}' (score: {score:.2f})"
                    )
                    
                    candidates.append(candidate)
        
        wb.close()
        
    except Exception as e:
        print(f"Error finding metric labels in {sheet_name}: {str(e)}")
    
    return candidates


def _find_best_metric_match(
    label_text: str,
    source_metrics: List[SourceMetric],
    min_fuzzy_score: float = 75.0
) -> Optional[Tuple[SourceMetric, str, float]]:
    """
    Find the best matching source metric for a label.
    
    Returns:
        Tuple of (SourceMetric, match_type, score) or None
    """
    label_lower = label_text.lower().strip()
    
    # Try exact ID match first
    for metric in source_metrics:
        if metric.metric_id.lower() == label_lower:
            return (metric, 'exact_id', 1.0)
    
    # Try exact name match
    for metric in source_metrics:
        if metric.metric_name.lower() == label_lower:
            return (metric, 'exact_name', 1.0)
    
    # Try fuzzy matching on name
    best_score = 0.0
    best_metric = None
    
    for metric in source_metrics:
        # Try different fuzzy matching algorithms
        score1 = fuzz.ratio(label_lower, metric.metric_name.lower())
        score2 = fuzz.partial_ratio(label_lower, metric.metric_name.lower())
        score3 = fuzz.token_sort_ratio(label_lower, metric.metric_name.lower())
        
        score = max(score1, score2, score3)
        
        if score > best_score:
            best_score = score
            best_metric = metric
    
    if best_score >= min_fuzzy_score:
        return (best_metric, 'fuzzy_name', best_score / 100.0)
    
    return None
