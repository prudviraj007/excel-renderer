"""Scoring and mapping proposal logic."""
from typing import List, Dict, Set, Tuple
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .models import (
    MappingCandidate, MonthHeaderCandidate, MetricLabelCandidate,
    SourceMetric
)
from .workbook_inspector import get_cell_value, is_cell_merged, WorkbookInfo


def propose_mappings(
    workbook_path: str,
    workbook_info: WorkbookInfo,
    source_metrics: List[SourceMetric],
    month_headers: List[MonthHeaderCandidate],
    metric_labels: List[MetricLabelCandidate],
    target_month: str,
    min_confidence: float = 0.5
) -> List[MappingCandidate]:
    """
    Propose mappings from source metrics to target cells.
    
    Args:
        workbook_path: Path to workbook
        workbook_info: Workbook metadata
        source_metrics: List of source metrics
        month_headers: List of month header candidates
        metric_labels: List of metric label candidates
        target_month: Target month (YYYY-MM)
        min_confidence: Minimum confidence threshold
    
    Returns:
        List of MappingCandidate objects
    """
    candidates = []
    
    # Filter month headers for target month
    target_month_headers = [h for h in month_headers if h.normalized_month == target_month]
    
    if not target_month_headers:
        print(f"No month headers found for target month {target_month}")
        return []
    
    # Group by sheet
    headers_by_sheet = defaultdict(list)
    for header in target_month_headers:
        headers_by_sheet[header.sheet_name].append(header)
    
    labels_by_sheet = defaultdict(list)
    for label in metric_labels:
        labels_by_sheet[label.sheet_name].append(label)
    
    # Process each sheet
    for sheet_name in set(headers_by_sheet.keys()) | set(labels_by_sheet.keys()):
        sheet_headers = headers_by_sheet.get(sheet_name, [])
        sheet_labels = labels_by_sheet.get(sheet_name, [])
        
        if not sheet_headers or not sheet_labels:
            continue
        
        # Propose mappings for this sheet
        sheet_candidates = _propose_sheet_mappings(
            workbook_path=workbook_path,
            workbook_info=workbook_info,
            sheet_name=sheet_name,
            source_metrics=source_metrics,
            month_headers=sheet_headers,
            metric_labels=sheet_labels,
            target_month=target_month
        )
        
        candidates.extend(sheet_candidates)
    
    # Filter by confidence
    candidates = [c for c in candidates if c.confidence >= min_confidence]
    
    # Mark ambiguous candidates
    candidates = _mark_ambiguous_mappings(candidates)
    
    return candidates


def _propose_sheet_mappings(
    workbook_path: str,
    workbook_info: WorkbookInfo,
    sheet_name: str,
    source_metrics: List[SourceMetric],
    month_headers: List[MonthHeaderCandidate],
    metric_labels: List[MetricLabelCandidate],
    target_month: str
) -> List[MappingCandidate]:
    """Propose mappings for a single sheet."""
    candidates = []
    
    try:
        wb = load_workbook(workbook_path, read_only=False, data_only=False)
        sheet = wb[sheet_name]
        
        # For each source metric that has a label match
        for source_metric in source_metrics:
            if source_metric.reporting_month != target_month:
                continue
            
            # Find labels that match this metric
            matching_labels = [
                label for label in metric_labels
                if label.matched_metric_id == source_metric.metric_id
            ]
            
            if not matching_labels:
                continue
            
            # For each matching label, try to find corresponding month header
            for label in matching_labels:
                for header in month_headers:
                    # Check if header and label are in compatible layout
                    if not _are_compatible(label, header):
                        continue
                    
                    # Calculate target cell location
                    target_cell = _calculate_target_cell(label, header)
                    
                    if not target_cell:
                        continue
                    
                    row, col = target_cell

                    # Reject intersections that land on the label or header anchor
                    if row == label.row and col == label.col:
                        continue
                    if row == header.row and col == header.col:
                        continue
                    if header.orientation == "horizontal" and col == label.col:
                        continue
                    if header.orientation == "vertical" and row == label.row:
                        continue
                    
                    # Get existing cell info
                    existing_value, is_formula = get_cell_value(sheet, row, col)
                    is_merged = is_cell_merged(sheet, row, col)
                    
                    # Calculate confidence
                    confidence = _calculate_mapping_confidence(
                        label=label,
                        header=header,
                        existing_value=existing_value,
                        is_formula=is_formula,
                        is_merged=is_merged
                    )
                    
                    # Build evidence
                    evidence = _build_mapping_evidence(
                        label=label,
                        header=header,
                        source_metric=source_metric
                    )
                    
                    # Determine state
                    state = 'suggested'
                    if is_formula:
                        state = 'blocked'  # Don't auto-approve formula cells
                    
                    candidate = MappingCandidate(
                        workbook_path=workbook_path,
                        workbook_name=workbook_info.name,
                        workbook_fingerprint=workbook_info.fingerprint,
                        sheet_name=sheet_name,
                        metric_id=source_metric.metric_id,
                        metric_name=source_metric.metric_name,
                        source_value=source_metric.value,
                        source_unit=source_metric.unit,
                        target_month=target_month,
                        orientation=header.orientation,
                        metric_label_cell=label.cell_address,
                        month_header_cell=header.cell_address,
                        target_value_cell=f"{get_column_letter(col)}{row}",
                        confidence=confidence,
                        match_evidence=evidence,
                        existing_value=existing_value,
                        has_formula=is_formula,
                        is_merged=is_merged,
                        state=state
                    )
                    
                    candidates.append(candidate)
        
        wb.close()
        
    except Exception as e:
        print(f"Error proposing mappings for {sheet_name}: {str(e)}")
    
    return candidates


def _are_compatible(label: MetricLabelCandidate, header: MonthHeaderCandidate) -> bool:
    """Check if a label and header are in compatible layout."""
    if header.orientation == "horizontal":
        # Month columns across a header row; metric names in rows to the left
        if label.row == header.row:
            return False
        if label.col >= header.col:
            return False
        return True
    # Vertical: month labels down a column; metric names in columns on a header row
    if label.col == header.col:
        return False
    if label.row >= header.row:
        return False
    return True


def _calculate_target_cell(
    label: MetricLabelCandidate,
    header: MonthHeaderCandidate
) -> Tuple[int, int]:
    """
    Calculate the target cell coordinates given label and header.
    
    Returns:
        Tuple of (row, col)
    """
    if header.orientation == 'horizontal':
        # Month in column, metric in row
        # Target is at (label_row, header_col)
        return (label.row, header.col)
    else:
        # Month in row, metric in column
        # Target is at (header_row, label_col)
        return (header.row, label.col)


def _calculate_mapping_confidence(
    label: MetricLabelCandidate,
    header: MonthHeaderCandidate,
    existing_value: any,
    is_formula: bool,
    is_merged: bool
) -> float:
    """Calculate confidence score for a mapping."""
    confidence = 0.0
    
    # Start with label match score
    confidence += label.match_score * 0.4
    
    # Add month header confidence
    confidence += header.confidence * 0.3
    
    # Bonus for exact matches
    if label.match_type in ['exact_id', 'exact_name']:
        confidence += 0.2
    
    # Penalty for formula cells
    if is_formula:
        confidence -= 0.3
    
    # Penalty for merged cells (not anchor)
    if is_merged:
        confidence -= 0.1
    
    # Slight bonus if cell is empty (clean target)
    if existing_value is None or existing_value == "":
        confidence += 0.1
    
    return max(0.0, min(1.0, confidence))


def _build_mapping_evidence(
    label: MetricLabelCandidate,
    header: MonthHeaderCandidate,
    source_metric: SourceMetric
) -> str:
    """Build evidence string for mapping."""
    parts = []
    
    parts.append(f"Metric: {label.match_type} match '{label.label_text}' at {label.cell_address}")
    parts.append(f"Month: {header.normalized_month} at {header.cell_address}")
    parts.append(f"Layout: {header.orientation}")
    parts.append(f"Source: {source_metric.metric_id} = {source_metric.value} {source_metric.unit}")
    
    return " | ".join(parts)


def _mark_ambiguous_mappings(candidates: List[MappingCandidate]) -> List[MappingCandidate]:
    """
    Mark mappings as ambiguous if there are multiple candidates for same metric/month.
    """
    # Group by (workbook, sheet, metric_id, target_month)
    groups = defaultdict(list)
    
    for candidate in candidates:
        key = (
            candidate.workbook_path,
            candidate.sheet_name,
            candidate.metric_id,
            candidate.target_month
        )
        groups[key].append(candidate)
    
    # Mark ambiguous if multiple candidates
    for key, group in groups.items():
        if len(group) > 1:
            # Check if they're actually different target cells
            unique_targets = set(c.target_value_cell for c in group)
            
            if len(unique_targets) > 1:
                # Multiple different target cells - ambiguous
                for candidate in group:
                    if candidate.state == 'suggested':
                        candidate.state = 'ambiguous'
    
    return candidates


def find_candidate_regions(
    workbook_path: str,
    sheet_name: str,
    month_headers: List[MonthHeaderCandidate],
    metric_labels: List[MetricLabelCandidate]
) -> List[Dict]:
    """
    Identify contiguous regions containing month headers and metric labels.
    
    Returns:
        List of region descriptors
    """
    regions = []
    
    # Group headers by orientation
    horizontal_headers = [h for h in month_headers if h.orientation == 'horizontal']
    vertical_headers = [h for h in month_headers if h.orientation == 'vertical']
    
    # Find horizontal regions (months in columns, metrics in rows)
    if horizontal_headers:
        # Group headers by row
        by_row = defaultdict(list)
        for h in horizontal_headers:
            by_row[h.row].append(h)
        
        for row, headers in by_row.items():
            headers_sorted = sorted(headers, key=lambda h: h.col)
            
            # Find labels in nearby rows
            nearby_labels = [
                l for l in metric_labels
                if abs(l.row - row) <= 20 and l.col < min(h.col for h in headers)
            ]
            
            if nearby_labels:
                region = {
                    'sheet': sheet_name,
                    'orientation': 'horizontal',
                    'header_row': row,
                    'header_cols': [h.col for h in headers_sorted],
                    'label_col': nearby_labels[0].col if nearby_labels else None,
                    'label_rows': [l.row for l in nearby_labels],
                    'months': [h.normalized_month for h in headers_sorted]
                }
                regions.append(region)
    
    # Find vertical regions (months in rows, metrics in columns)
    if vertical_headers:
        # Group headers by column
        by_col = defaultdict(list)
        for h in vertical_headers:
            by_col[h.col].append(h)
        
        for col, headers in by_col.items():
            headers_sorted = sorted(headers, key=lambda h: h.row)
            
            # Find labels in nearby columns
            nearby_labels = [
                l for l in metric_labels
                if abs(l.col - col) <= 10 and l.row < min(h.row for h in headers)
            ]
            
            if nearby_labels:
                region = {
                    'sheet': sheet_name,
                    'orientation': 'vertical',
                    'header_col': col,
                    'header_rows': [h.row for h in headers_sorted],
                    'label_row': nearby_labels[0].row if nearby_labels else None,
                    'label_cols': [l.col for l in nearby_labels],
                    'months': [h.normalized_month for h in headers_sorted]
                }
                regions.append(region)
    
    return regions
