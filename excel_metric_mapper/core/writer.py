"""Workbook writer - safely writes values to Excel files."""
import shutil
import tempfile
from pathlib import Path
from typing import List, Tuple
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from .models import ProposedChange, ExecutionResult, MappingCandidate


def write_changes(
    workbook_path: str,
    changes: List[ProposedChange],
    output_dir: str,
    dry_run: bool = False
) -> ExecutionResult:
    """
    Write approved changes to a workbook copy.
    
    Args:
        workbook_path: Path to source workbook
        changes: List of ProposedChange objects
        output_dir: Directory for output files
        dry_run: If True, don't actually write
    
    Returns:
        ExecutionResult object
    """
    start_time = datetime.now()
    workbook_name = Path(workbook_path).name
    output_path = Path(output_dir) / f"updated_{workbook_name}"
    
    result = ExecutionResult(
        workbook_path=workbook_path,
        output_path=str(output_path),
        success=False,
        changes_applied=0,
        changes_skipped=0
    )
    
    if dry_run:
        result.warnings.append("DRY RUN - no changes written")
        result.success = True
        return result
    
    try:
        # Ensure output directory exists
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        
        # Load workbook (preserve VBA if .xlsm)
        keep_vba = workbook_path.endswith('.xlsm')
        wb = load_workbook(workbook_path, keep_vba=keep_vba, data_only=False)
        
        # Group changes by sheet
        changes_by_sheet = {}
        for change in changes:
            if change.sheet_name not in changes_by_sheet:
                changes_by_sheet[change.sheet_name] = []
            changes_by_sheet[change.sheet_name].append(change)
        
        # Apply changes
        for sheet_name, sheet_changes in changes_by_sheet.items():
            if sheet_name not in wb.sheetnames:
                result.errors.append(f"Sheet not found: {sheet_name}")
                result.changes_skipped += len(sheet_changes)
                continue
            
            sheet = wb[sheet_name]
            
            for change in sheet_changes:
                try:
                    # Parse cell address
                    row, col = _parse_cell_address(change.target_cell)
                    
                    # Safety checks
                    if change.has_formula:
                        result.warnings.append(
                            f"Skipped {change.target_cell} in {sheet_name}: contains formula"
                        )
                        result.changes_skipped += 1
                        continue
                    
                    if change.is_merged:
                        result.warnings.append(
                            f"Warning: {change.target_cell} in {sheet_name} is merged"
                        )
                    
                    # Write value
                    cell = sheet.cell(row=row, column=col)
                    old_value = cell.value
                    cell.value = change.proposed_value
                    
                    result.changes_applied += 1
                    
                except Exception as e:
                    result.errors.append(
                        f"Failed to write {change.target_cell} in {sheet_name}: {str(e)}"
                    )
                    result.changes_skipped += 1
        
        # Save to temporary file first
        with tempfile.NamedTemporaryFile(delete=False, suffix=Path(workbook_path).suffix) as tmp:
            temp_path = tmp.name
        
        wb.save(temp_path)
        wb.close()
        
        # Move to final location
        shutil.move(temp_path, output_path)
        
        # Verify the output
        verification_result = verify_output(str(output_path), changes)
        if not verification_result[0]:
            result.warnings.append(f"Verification warning: {verification_result[1]}")
        
        result.success = True
        
    except Exception as e:
        result.errors.append(f"Failed to write workbook: {str(e)}")
        result.success = False
    
    finally:
        end_time = datetime.now()
        result.execution_time = (end_time - start_time).total_seconds()
        result.timestamp = end_time
    
    return result


def verify_output(
    output_path: str,
    changes: List[ProposedChange]
) -> Tuple[bool, str]:
    """
    Verify that changes were written correctly.
    
    Returns:
        Tuple of (success, message)
    """
    try:
        wb = load_workbook(output_path, read_only=True, data_only=True)
        
        verification_failed = []
        
        # Group changes by sheet
        changes_by_sheet = {}
        for change in changes:
            if change.sheet_name not in changes_by_sheet:
                changes_by_sheet[change.sheet_name] = []
            changes_by_sheet[change.sheet_name].append(change)
        
        for sheet_name, sheet_changes in changes_by_sheet.items():
            if sheet_name not in wb.sheetnames:
                verification_failed.append(f"Sheet {sheet_name} not found")
                continue
            
            sheet = wb[sheet_name]
            
            for change in sheet_changes:
                if change.has_formula:
                    continue  # Skip formula cells
                
                row, col = _parse_cell_address(change.target_cell)
                cell = sheet.cell(row=row, column=col)
                actual_value = cell.value
                
                # Compare values (with some tolerance for floats)
                if isinstance(change.proposed_value, float) and isinstance(actual_value, (int, float)):
                    if abs(change.proposed_value - float(actual_value)) > 0.001:
                        verification_failed.append(
                            f"{sheet_name}!{change.target_cell}: expected {change.proposed_value}, got {actual_value}"
                        )
                elif actual_value != change.proposed_value:
                    verification_failed.append(
                        f"{sheet_name}!{change.target_cell}: expected {change.proposed_value}, got {actual_value}"
                    )
        
        wb.close()
        
        if verification_failed:
            return False, "; ".join(verification_failed)
        else:
            return True, "All changes verified"
    
    except Exception as e:
        return False, f"Verification error: {str(e)}"


def prepare_proposed_changes(
    candidates: List[MappingCandidate]
) -> List[ProposedChange]:
    """
    Convert approved MappingCandidates to ProposedChange objects.
    
    Args:
        candidates: List of approved MappingCandidate objects
    
    Returns:
        List of ProposedChange objects
    """
    changes = []
    
    for candidate in candidates:
        if candidate.state != 'approved':
            continue
        
        warnings = []
        
        if candidate.has_formula:
            warnings.append("Target cell contains a formula")
        
        if candidate.is_merged:
            warnings.append("Target cell is in a merged range")
        
        if candidate.existing_value is not None and candidate.existing_value != "":
            warnings.append(f"Will overwrite existing value: {candidate.existing_value}")
        
        change = ProposedChange(
            workbook_path=candidate.workbook_path,
            sheet_name=candidate.sheet_name,
            target_cell=candidate.target_value_cell,
            metric_id=candidate.metric_id,
            metric_name=candidate.metric_name,
            source_value=candidate.source_value,
            source_unit=candidate.source_unit,
            existing_value=candidate.existing_value,
            existing_formula=None,  # TODO: extract formula if needed
            has_formula=candidate.has_formula,
            is_merged=candidate.is_merged,
            proposed_value=candidate.source_value,
            confidence=candidate.confidence,
            warnings=warnings
        )
        
        changes.append(change)
    
    return changes


def _parse_cell_address(address: str) -> Tuple[int, int]:
    """
    Parse cell address like 'B5' to (row, col) tuple.
    
    Returns:
        Tuple of (row, col) in 1-indexed format
    """
    import re
    match = re.match(r'^([A-Z]+)(\d+)$', address.upper())
    if not match:
        raise ValueError(f"Invalid cell address: {address}")
    
    col_str, row_str = match.groups()
    col = column_index_from_string(col_str)
    row = int(row_str)
    
    return row, col


def create_backup(workbook_path: str, backup_dir: str) -> str:
    """
    Create a backup copy of a workbook.
    
    Returns:
        Path to backup file
    """
    Path(backup_dir).mkdir(parents=True, exist_ok=True)
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    workbook_name = Path(workbook_path).stem
    extension = Path(workbook_path).suffix
    backup_name = f"{workbook_name}_backup_{timestamp}{extension}"
    backup_path = Path(backup_dir) / backup_name
    
    shutil.copy2(workbook_path, backup_path)
    
    return str(backup_path)
