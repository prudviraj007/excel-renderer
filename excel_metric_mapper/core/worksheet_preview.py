"""Build a small worksheet region preview around a mapping candidate."""
from __future__ import annotations

import re
from typing import Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .models import MappingCandidate


def parse_cell_address(address: str) -> Tuple[int, int]:
    match = re.match(r"^([A-Za-z]+)(\d+)$", address.strip())
    if not match:
        raise ValueError(f"Invalid cell address: {address}")
    col = column_index_from_string(match.group(1).upper())
    row = int(match.group(2))
    return row, col


def build_preview_dataframe(
    candidate: MappingCandidate,
    radius: int = 4,
) -> pd.DataFrame:
    wb = load_workbook(candidate.workbook_path, read_only=True, data_only=True)
    try:
        sheet = wb[candidate.sheet_name]
        target_row, target_col = parse_cell_address(candidate.target_value_cell)

        start_row = max(1, target_row - radius)
        end_row = min(sheet.max_row or target_row, target_row + radius)
        start_col = max(1, target_col - radius)
        end_col = min(sheet.max_column or target_col, target_col + radius)

        index_labels = []
        rows = []
        styles: dict[tuple[int, str], str] = {}

        for row in range(start_row, end_row + 1):
            index_labels.append(str(row))
            row_values = []
            for col in range(start_col, end_col + 1):
                addr = f"{get_column_letter(col)}{row}"
                cell = sheet.cell(row=row, column=col)
                raw = cell.value
                display = "" if raw is None else str(raw)[:24]

                if addr == candidate.target_value_cell:
                    display = f"▶ {display or '(empty)'}"
                    styles[(len(rows), get_column_letter(col))] = "target"
                elif addr == candidate.month_header_cell:
                    display = f"📅 {display}"
                    styles[(len(rows), get_column_letter(col))] = "month"
                elif addr == candidate.metric_label_cell:
                    display = f"🏷 {display}"
                    styles[(len(rows), get_column_letter(col))] = "label"

                row_values.append(display)
            rows.append(row_values)

        columns = [get_column_letter(c) for c in range(start_col, end_col + 1)]
        df = pd.DataFrame(rows, index=index_labels, columns=columns)
        df.index.name = "Row"
        df.attrs["cell_styles"] = styles
        return df
    finally:
        wb.close()


def style_preview(df: pd.DataFrame):
    styles = df.attrs.get("cell_styles", {})

    def _apply(data: pd.DataFrame) -> pd.DataFrame:
        out = pd.DataFrame("", index=data.index, columns=data.columns)
        for (row_idx, col_name), kind in styles.items():
            if row_idx < len(data.index) and col_name in data.columns:
                if kind == "target":
                    out.iloc[row_idx, data.columns.get_loc(col_name)] = (
                        "background-color: #ffe082; font-weight: bold"
                    )
                elif kind == "month":
                    out.iloc[row_idx, data.columns.get_loc(col_name)] = (
                        "background-color: #bbdefb"
                    )
                elif kind == "label":
                    out.iloc[row_idx, data.columns.get_loc(col_name)] = (
                        "background-color: #c8e6c9"
                    )
        return out

    return df.style.apply(_apply, axis=None)
